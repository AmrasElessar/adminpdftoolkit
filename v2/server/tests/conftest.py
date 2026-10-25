from __future__ import annotations

import io
import time
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from fastapi.testclient import TestClient

from adminpdf.app import create_app
from adminpdf.config import Settings

LOCAL = ("127.0.0.1", 50000)


def make_pdf(pages: int = 3, *, text: str = "Sayfa", size: tuple[float, float] = (595, 842)) -> bytes:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=size[0], height=size[1])
        page.insert_text((72, 100), f"{text} {i + 1}", fontsize=24)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_image_pdf() -> bytes:
    """A page carrying a large, high-DPI photographic-ish image."""
    from PIL import Image

    img = Image.effect_mandelbrot((2400, 2400), (-2, -1.5, 1, 1.5), 100).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(36, 36, 336, 336), stream=buf.getvalue())
    data: bytes = doc.tobytes(deflate=True)
    doc.close()
    return data


def make_png(w: int = 200, h: int = 100) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(data_dir=tmp_path / "data", web_dist=tmp_path / "nodist", max_upload_mb=5, job_timeout_seconds=60)


class Api:
    """Thin client that does the CSRF handshake and waits for jobs."""

    def __init__(self, client: TestClient) -> None:
        self.c = client
        token = client.get("/api/session").json()["csrf"]
        client.headers["X-CSRF-Token"] = token

    def upload(self, data: bytes, name: str = "test.pdf") -> dict[str, Any]:
        r = self.c.put("/api/files", content=data, headers={"X-Filename": name})
        assert r.status_code == 201, r.text
        result: dict[str, Any] = r.json()
        return result

    def start(self, op: str, file_ids: list[str], params: dict[str, Any] | None = None, **kw: Any) -> Any:
        inputs = [{"id": i, **kw} for i in file_ids]
        return self.c.post(f"/api/ops/{op}", json={"inputs": inputs, "params": params or {}})

    def wait(self, job_id: str, timeout: float = 60) -> dict[str, Any]:
        deadline = time.time() + timeout
        while time.time() < deadline:
            job: dict[str, Any] = self.c.get(f"/api/jobs/{job_id}").json()
            if job["status"] in ("done", "failed", "cancelled"):
                return job
            time.sleep(0.05)
        raise AssertionError("job did not finish")

    def run(self, op: str, file_ids: list[str], params: dict[str, Any] | None = None, **kw: Any) -> dict[str, Any]:
        r = self.start(op, file_ids, params, **kw)
        assert r.status_code == 202, r.text
        return self.wait(r.json()["id"])

    def download(self, file_id: str) -> bytes:
        r = self.c.get(f"/api/files/{file_id}")
        assert r.status_code == 200
        return r.content


@pytest.fixture
def client(settings: Settings) -> Iterator[TestClient]:
    with TestClient(create_app(settings), base_url="http://localhost", client=LOCAL) as c:
        yield c


@pytest.fixture
def api(client: TestClient) -> Api:
    return Api(client)


def make_docx(text: str = "Merhaba dünya — İĞŞÇÖÜ") -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Başlık", level=1)
    p = d.add_paragraph()
    p.add_run(text).bold = True
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Ad"
    t.cell(0, 1).text = "Tutar"
    t.cell(1, 0).text = "Şükrü"
    t.cell(1, 1).text = "1.250,50"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Satış"
    ws.append(["Ürün", "Adet"])
    ws.append(["Çay", 12])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_table_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    rows = [["Ad", "Tutar", "Not"], ["Ali", "1.250,50", "=HYPERLINK(1)"], ["Ayşe", "300", "ok"]]
    x0, y0, w, h = 72, 72, 150, 24
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 4, rect.y1 - 7), value, fontsize=11)
    data: bytes = doc.tobytes()
    doc.close()
    return data
