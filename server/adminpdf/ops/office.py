"""Office conversions: PDF → Word / Excel, and Word / Excel / PowerPoint → PDF.

* PDF → Word: ``pdf2docx`` (MIT) rebuilds paragraphs, tables and images.
* PDF → Excel: PyMuPDF's table finder; one sheet per table.
* Office → PDF: LibreOffice in headless mode when installed (best fidelity),
  otherwise a pure-Python fallback for .docx / .xlsx rendered through
  PyMuPDF's HTML engine.
"""

from __future__ import annotations

import html
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Literal

import pymupdf

from ..core import proc
from ..core.errors import UserError
from ..core.pdf import open_pdf, parse_page_spec
from .base import Context, Field, InputFile, Operation, Params, register

# OOXML containers are zips: refuse decompression bombs before parsing.
MAX_OOXML_UNCOMPRESSED = 512 * 1024 * 1024
MAX_OOXML_RATIO = 200
MAX_SHEET_CELLS = 2_000_000
SOFFICE_TIMEOUT = 300

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def check_ooxml(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as e:
        raise UserError(
            "Dosya bozuk veya geçerli bir Office belgesi değil.", "The file is not a valid Office document."
        ) from e
    total = sum(i.file_size for i in infos)
    packed = sum(i.compress_size for i in infos) or 1
    if total > MAX_OOXML_UNCOMPRESSED or total / packed > MAX_OOXML_RATIO:
        raise UserError(
            "Belge açıldığında çok büyük (zip bombası?).", "The document expands to an unsafe size (zip bomb?)."
        )


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    ):
        if candidate.is_file():
            return str(candidate)
    return None


# -- PDF → Word ----------------------------------------------------------------
class ToWordParams(Params):
    pages: str = Field("", tr="Sayfalar (boş = tümü)", en="Pages (empty = all)", max_length=2000)


@register
class PdfToWord(Operation):
    name = "pdf_to_word"
    category = "convert"
    title_tr, title_en = "PDF'ten Word'e", "PDF to Word"
    description_tr = "PDF'i düzenlenebilir Word (.docx) belgesine dönüştürür; paragraf, tablo ve görselleri korur."
    description_en = "Convert a PDF into an editable Word (.docx) document, keeping paragraphs, tables and images."
    Params = ToWordParams

    def run(self, ctx: Context, params: ToWordParams) -> None:
        from pdf2docx import Converter

        inp = ctx.inputs[0]
        with open_pdf(inp.path, inp.password) as doc:
            pages = sorted(set(parse_page_spec(params.pages, doc.page_count)))
        out = ctx.output(f"{inp.stem}.docx")
        ctx.progress(0.05, "parse")
        cv = Converter(str(inp.path), password=inp.password or None)
        try:
            cv.convert(str(out), pages=pages, multi_processing=False)
        finally:
            cv.close()


# -- PDF → Excel ---------------------------------------------------------------
class ToExcelParams(Params):
    layout: Literal["sheet_per_table", "single_sheet"] = Field(
        "sheet_per_table",
        tr="Yerleşim",
        en="Layout",
        ui={
            "options": {
                "sheet_per_table": {"tr": "Her tablo ayrı sayfa", "en": "One sheet per table"},
                "single_sheet": {"tr": "Tüm tablolar tek sayfada", "en": "All tables on one sheet"},
            }
        },
    )
    numbers: Literal["tr", "en", "text"] = Field(
        "tr",
        tr="Sayı biçimi",
        en="Number format",
        help_tr="Hücrelerdeki sayılar Excel sayısına çevrilsin mi?",
        help_en="Turn numeric-looking cells into real Excel numbers?",
        ui={
            "options": {
                "tr": {"tr": "1.234,56 (Türkçe)", "en": "1.234,56 (Turkish)"},
                "en": {"tr": "1,234.56 (İngilizce)", "en": "1,234.56 (English)"},
                "text": {"tr": "Metin olarak bırak", "en": "Keep as text"},
            }
        },
    )
    pages: str = Field("", tr="Sayfalar (boş = tümü)", en="Pages (empty = all)", max_length=2000)


_NUM = {
    "tr": re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+(,\d+)?$"),
    "en": re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$"),
}


def to_cell(value: Any, numbers: str) -> Any:
    if value is None:
        return None
    text = _ILLEGAL_XML.sub("", str(value)).strip()
    if numbers in _NUM and _NUM[numbers].match(text):
        thousands, decimal = (".", ",") if numbers == "tr" else (",", ".")
        num = text.replace(thousands, "").replace(decimal, ".")
        try:
            f = float(num)
            return int(f) if "." not in num and abs(f) < 2**53 else f
        except ValueError:
            pass
    return text


def write_cell(ws: Any, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        # Never let PDF text become a live formula (CSV/Excel injection).
        cell.data_type = "s"


@register
class PdfToExcel(Operation):
    name = "pdf_to_excel"
    category = "convert"
    title_tr, title_en = "PDF'ten Excel'e", "PDF to Excel"
    description_tr = "PDF'teki tabloları algılayıp Excel (.xlsx) çalışma kitabına aktarır."
    description_en = "Detect the tables in a PDF and export them to an Excel (.xlsx) workbook."
    Params = ToExcelParams

    def run(self, ctx: Context, params: ToExcelParams) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        inp = ctx.inputs[0]
        wb = Workbook()
        wb.remove(wb.active)
        single = wb.create_sheet("Tablolar") if params.layout == "single_sheet" else None
        row_cursor = 1
        found = 0
        cells = 0
        with open_pdf(inp.path, inp.password) as doc:
            targets = sorted(set(parse_page_spec(params.pages, doc.page_count)))
            for k, i in enumerate(targets):
                for t_idx, table in enumerate(doc[i].find_tables().tables, start=1):
                    rows = table.extract()
                    if not rows:
                        continue
                    found += 1
                    cells += sum(len(r) for r in rows)
                    if cells > MAX_SHEET_CELLS:
                        raise UserError("Tablo verisi çok büyük.", "Too much table data.")
                    if single is not None:
                        ws, start = single, row_cursor
                        ws.cell(row=start, column=1, value=f"Sayfa {i + 1} — Tablo {t_idx}").font = Font(bold=True)
                        start += 1
                    else:
                        ws, start = wb.create_sheet(f"S{i + 1}-T{t_idx}"[:31]), 1
                    for r_off, row in enumerate(rows):
                        for c_off, value in enumerate(row):
                            write_cell(ws, start + r_off, c_off + 1, to_cell(value, params.numbers))
                    for c in range(1, len(rows[0]) + 1):
                        ws.cell(row=start, column=c).font = Font(bold=True)
                    row_cursor = start + len(rows) + 1
                ctx.progress((k + 1) / len(targets), "")
        if not found:
            raise UserError(
                "PDF'te tablo bulunamadı. Taranmış bir belgeyse önce OCR uygulayın.",
                "No tables found. If the document is scanned, run OCR first.",
            )
        wb.save(ctx.output(f"{inp.stem}.xlsx"))


# -- Office → PDF --------------------------------------------------------------
class OfficeParams(Params):
    engine: Literal["auto", "libreoffice", "builtin"] = Field(
        "auto",
        tr="Motor",
        en="Engine",
        help_tr="LibreOffice en iyi sonucu verir; yüklü değilse dahili motor kullanılır.",
        help_en="LibreOffice gives the best fidelity; the built-in engine is used when it is not installed.",
        ui={
            "options": {
                "auto": {"tr": "Otomatik", "en": "Automatic"},
                "libreoffice": {"tr": "LibreOffice", "en": "LibreOffice"},
                "builtin": {"tr": "Dahili (basit)", "en": "Built-in (basic)"},
            }
        },
    )


# A throwaway LibreOffice profile that treats every document as hostile: no
# macros, no active content (OLE / DDE), links and external data never
# updated, formulas never recalculated on load, and every network request
# sent to a proxy that does not exist (port 9, "discard"), so a crafted
# document can neither read local files nor reach the network (SSRF).
# Unknown keys are ignored by older LibreOffice versions.
_LO_HARDENING = {
    "/org.openoffice.Office.Common/Security/Scripting": {
        "MacroSecurityLevel": ("int", "3"),
        "DisableMacrosExecution": ("boolean", "true"),
        "DisableActiveContent": ("boolean", "true"),
        "BlockUntrustedRefererLinks": ("boolean", "true"),
    },
    # Calc: 0 always, 1 never, 2 on request. Writer: 0 never, 1 on request.
    "/org.openoffice.Office.Calc/Content/Update": {"Link": ("int", "1")},
    "/org.openoffice.Office.Writer/Content/Update": {"Link": ("int", "0")},
    "/org.openoffice.Office.Calc/Formula/Load": {
        "OOXMLRecalcMode": ("int", "1"),
        "ODFRecalcMode": ("int", "1"),
    },
    "/org.openoffice.Inet/Settings": {
        "ooInetProxyType": ("int", "2"),
        "ooInetHTTPProxyName": ("string", "127.0.0.1"),
        "ooInetHTTPProxyPort": ("int", "9"),
        "ooInetHTTPSProxyName": ("string", "127.0.0.1"),
        "ooInetHTTPSProxyPort": ("int", "9"),
        "ooInetFTPProxyName": ("string", "127.0.0.1"),
        "ooInetFTPProxyPort": ("int", "9"),
        "ooInetNoProxy": ("string", ""),
    },
}


def write_hardened_profile(profile: Path) -> None:
    items = []
    for path, props in _LO_HARDENING.items():
        for name, (kind, value) in props.items():
            items.append(
                f'<item oor:path="{path}"><prop oor:name="{name}" oor:op="fuse" oor:type="xs:{kind}">'
                f"<value>{html.escape(value)}</value></prop></item>"
            )
    (profile / "user").mkdir(parents=True)
    (profile / "user" / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
        'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n' + "\n".join(items) + "\n</oor:items>\n",
        encoding="utf-8",
    )


def libreoffice_to_pdf(soffice: str, inp: InputFile, out: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="apt-lo-") as tmp:
        tmpdir = Path(tmp)
        src = tmpdir / f"input{inp.path.suffix or Path(inp.name).suffix}"
        shutil.copyfile(inp.path, src)
        write_hardened_profile(tmpdir / "profile")
        profile = (tmpdir / "profile").as_uri()
        result = proc.run(
            [
                soffice,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--norestore",
                "--nolockcheck",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmpdir / "out"),
                str(src),
            ],
            timeout=SOFFICE_TIMEOUT,
        )
        produced = tmpdir / "out" / "input.pdf"
        if result.returncode != 0 or not produced.is_file():
            raise UserError("LibreOffice belgeyi dönüştüremedi.", "LibreOffice could not convert the document.")
        shutil.move(produced, out)


def _esc(value: Any) -> str:
    return html.escape(_ILLEGAL_XML.sub("", str(value)))


def docx_to_html(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    parts: list[str] = []

    def runs(p: Paragraph) -> str:
        out = []
        for r in p.runs:
            t = _esc(r.text)
            if r.bold:
                t = f"<b>{t}</b>"
            if r.italic:
                t = f"<i>{t}</i>"
            if r.underline:
                t = f"<u>{t}</u>"
            out.append(t)
        return "".join(out) or "&nbsp;"

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            style = (block.style.name if block.style is not None else "").lower()
            align = {1: "center", 2: "right", 3: "justify"}.get(int(block.alignment or 0), "left")
            m = re.match(r"heading (\d)", style)
            tag = f"h{min(int(m.group(1)), 6)}" if m else ("li" if "list" in style else "p")
            parts.append(f'<{tag} style="text-align:{align}">{runs(block)}</{tag}>')
        elif isinstance(block, Table):
            rows = "".join(
                "<tr>"
                + "".join(f"<td>{'<br>'.join(_esc(p.text) for p in c.paragraphs)}</td>" for c in r.cells)
                + "</tr>"
                for r in block.rows
            )
            parts.append(f"<table>{rows}</table>")
    return "".join(parts)


def xlsx_to_html(path: Path) -> str:
    from openpyxl import load_workbook

    # Stored uploads have no extension, which openpyxl insists on for paths;
    # a file object sidesteps the check.
    fh = path.open("rb")
    wb = load_workbook(fh, read_only=True, data_only=True)
    parts: list[str] = []
    cells = 0
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_SHEET_CELLS:
                    raise UserError("Çalışma kitabı çok büyük.", "The workbook is too large.")
                if any(v is not None for v in row):
                    rows.append("<tr>" + "".join(f"<td>{'' if v is None else _esc(v)}</td>" for v in row) + "</tr>")
            if rows:
                parts.append(f"<h2>{_esc(ws.title)}</h2><table>{''.join(rows)}</table>")
    finally:
        wb.close()
        fh.close()
    return "".join(parts)


_HTML_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.35; }
h1 { font-size: 20pt; } h2 { font-size: 15pt; } h3 { font-size: 13pt; }
table { border-collapse: collapse; margin: 6pt 0; }
td { border: 0.5pt solid #999; padding: 2pt 4pt; font-size: 9pt; vertical-align: top; }
"""


def html_to_pdf(markup: str, out: Path, *, landscape: bool = False) -> None:
    story = pymupdf.Story(html=markup or "<p></p>", user_css=_HTML_CSS)
    page = pymupdf.paper_rect("a4-l" if landscape else "a4")
    where = pymupdf.Rect(page.x0 + 48, page.y0 + 48, page.x1 - 48, page.y1 - 48)
    writer = pymupdf.DocumentWriter(str(out))
    more = True
    try:
        while more:
            dev = writer.begin_page(page)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
    finally:
        writer.close()


@register
class OfficeToPdf(Operation):
    name = "office_to_pdf"
    category = "convert"
    title_tr, title_en = "Office'ten PDF'e", "Office to PDF"
    description_tr = "Word (.docx), Excel (.xlsx) ve PowerPoint (.pptx) belgelerini PDF'e dönüştürür."
    description_en = "Convert Word (.docx), Excel (.xlsx) and PowerPoint (.pptx) documents to PDF."
    accepts = frozenset({"docx", "xlsx", "pptx"})
    max_inputs = 50
    Params = OfficeParams

    def run(self, ctx: Context, params: OfficeParams) -> None:
        soffice = find_soffice() if params.engine != "builtin" else None
        if params.engine == "libreoffice" and soffice is None:
            raise UserError("LibreOffice yüklü değil.", "LibreOffice is not installed.")
        for k, inp in enumerate(ctx.inputs):
            check_ooxml(inp.path)
            out = ctx.output(f"{inp.stem}.pdf")
            if soffice:
                libreoffice_to_pdf(soffice, inp, out)
            else:
                ext = Path(inp.name).suffix.lower()
                if ext == ".docx":
                    html_to_pdf(docx_to_html(inp.path), out)
                elif ext == ".xlsx":
                    html_to_pdf(xlsx_to_html(inp.path), out, landscape=True)
                else:
                    raise UserError(
                        "PowerPoint dönüştürmek için LibreOffice gerekir.",
                        "Converting PowerPoint needs LibreOffice.",
                    )
            ctx.progress((k + 1) / len(ctx.inputs), inp.name)
