from __future__ import annotations

import html
import io
import random
import zipfile
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from openpyxl import Workbook, load_workbook

from adminpdf.core.errors import UserError
from adminpdf.plugins.calllog.distribution import (
    check_weights,
    distribute,
    ratio_counts,
    sequential_counts,
    team_file_names,
)
from adminpdf.plugins.calllog.parsing import (
    CALL_LOG_COLUMNS,
    is_record_start,
    normalize_lines,
    parse_call_log_lines,
    parse_qa,
)
from adminpdf.plugins.calllog.tabular import (
    FilterRule,
    Table,
    apply_rules,
    deduplicate,
    normalize_phone,
    read_xlsx,
    safe_value,
    write_workbook,
)

from .conftest import Api, make_image_pdf, make_pdf, make_table_pdf

# -- synthetic PDFs ------------------------------------------------------------
Record = dict[str, str]


def record(no: int, name: str, phone: str, **answers: str) -> Record:
    return {"no": str(no), "name": name, "phone": phone, **answers}


def call_log_lines(rec: Record) -> list[str]:
    return [
        rec["no"],
        rec["name"],
        rec["phone"],
        rec.get("status", "ended"),
        rec.get("date", "01.02.2026 10:15"),
        rec.get("duration", "2 dk 35 sn"),
        "S: Ağrı / romatizma",
        f"C: {rec.get('agri', 'Var')}",
        "S: Termal / kaplıca",
        f"C: {rec.get('termal', 'Evet')}",
        "S: Yaş",
        f"C: {rec.get('yas', '45')}",
        "S: Medeni durum",
        "C: Evli",
        "S: Meslek",
        f"C: {rec.get('meslek', 'Öğretmen')}",
        "S: İkamet ili",
        f"C: {rec.get('il', 'Ankara')}",
    ]


def make_call_log_pdf(records: list[Record], per_page: int = 2) -> bytes:
    """Mimics the 360 export: a "Seçilen kayıtlar" banner, a MÜŞTERİ column
    header, then one text block per record with S:/C: Q&A lines."""
    doc = pymupdf.open()
    chunks = [records[i : i + per_page] for i in range(0, len(records), per_page)] or [[]]
    for n, chunk in enumerate(chunks):
        lines = [f"Seçilen kayıtlar ({len(records)})", "# MÜŞTERİ TELEFON DURUM TARİH SÜRE"] if n == 0 else []
        for rec in chunk:
            lines.extend(call_log_lines(rec))
        body = "".join(f"<p style='margin:0'>{html.escape(ln)}</p>" for ln in lines)
        page = doc.new_page()
        page.insert_htmlbox(pymupdf.Rect(36, 36, 559, 806), body, css="p{font-size:9px}")
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_grid_pdf(rows: list[list[str]], pages: int = 1) -> bytes:
    """A ruled table; the header row is repeated on every page."""
    doc = pymupdf.open()
    header, body = rows[0], rows[1:]
    per = max(1, -(-len(body) // pages))
    for p in range(pages):
        page = doc.new_page()
        chunk = [header, *body[p * per : (p + 1) * per]]
        for r, row in enumerate(chunk):
            for c, value in enumerate(row):
                rect = pymupdf.Rect(50 + c * 140, 50 + r * 22, 50 + (c + 1) * 140, 50 + (r + 1) * 22)
                page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                page.insert_text((rect.x0 + 4, rect.y1 - 7), value, fontsize=10)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_xlsx_rows(rows: list[list[Any]], title: str = "Kayıtlar") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheet_values(data: bytes, sheet: int | str = 0) -> list[list[Any]]:
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    return [[c.value for c in row] for row in ws.iter_rows()]


RECORDS = [
    record(1, "Ahmet Yılmaz", "+90 532 111 22 33"),
    record(2, "Ayşe Kaya", "0533 444 55 66", agri="Yok"),
    record(3, '=HYPERLINK("http://x")', "0544 777 88 99"),
]
RECORDS_2 = [
    record(1, "Ahmet Y.", "05321112233"),  # same phone as record 1 above
    record(2, "Şükrü Öz", "5559990011", il="İzmir"),
]


# -- parser units (ported from v1 tests/test_call_log_parser.py) ---------------
def test_parse_qa() -> None:
    assert parse_qa(["S: Ağrı / romatizma", "C: Var", "S: Yaş", "C: Otuz üç"]) == {
        "Ağrı / romatizma": "Var",
        "Yaş": "Otuz üç",
    }
    assert parse_qa(["S: Meslek", "C: Mühendislik", "ek satır", "S: İl", "C: Ankara"])["Meslek"] == (
        "Mühendislik ek satır"
    )
    assert parse_qa(["S: Sigara", "C:", "S: Yaş", "C: 40"]) == {"Sigara": "", "Yaş": "40"}


def test_normalize_lines_and_record_start() -> None:
    assert normalize_lines("satir1\n\n  satir2  \n") == ["satir1", "satir2"]
    assert normalize_lines("S: Yaş\\nC: 40") == ["S: Yaş", "C: 40"]
    assert is_record_start("1") and is_record_start("999")
    assert not any(is_record_start(x) for x in ("1.0", "a1", "", "²", "0"))


def test_parse_call_log_lines() -> None:
    lines = ["Seçilen kayıtlar (2)", *call_log_lines(RECORDS[0]), *call_log_lines(RECORDS_2[1])]
    recs = parse_call_log_lines(lines)
    assert [r["Kayıt No"] for r in recs] == ["1", "2"]
    assert recs[0]["Müşteri"] == "Ahmet Yılmaz"
    assert recs[0]["Telefon"] == "+90 532 111 22 33"
    assert recs[0]["Durum"] == "ended"
    assert recs[0]["Tarih"] == "01.02.2026 10:15"
    assert recs[0]["Süre"] == "2 dk 35 sn"
    assert recs[1]["İkamet ili"] == "İzmir"
    assert recs[1]["AI Özeti (Ham)"].startswith("S: Ağrı / romatizma\nC: Var")


# -- phone / dedup / filter ------------------------------------------------------
@pytest.mark.parametrize(
    "raw",
    ["+90 532 111 22 33", "0532 111 22 33", "05321112233", "5321112233", "0090 (532) 111-22-33", "90 532 1112233"],
)
def test_normalize_phone_turkish_formats(raw: str) -> None:
    assert normalize_phone(raw) == "5321112233"


def test_normalize_phone_other() -> None:
    assert normalize_phone("") == ""
    assert normalize_phone(None) == ""
    assert normalize_phone("+1 212 555 0100") == "12125550100"


def test_deduplicate_keeps_first_and_skips_empty_keys() -> None:
    t = Table(
        ["Sıra", "Telefon", "Ad"],
        [[1, "0532 111 22 33", "a"], [2, "", "b"], [3, "+905321112233", "c"], [4, "", "d"], [5, "5550000000", "e"]],
    )
    removed = deduplicate(t, ["telefon"])  # column match is case-insensitive
    assert removed == [2]
    assert [r[2] for r in t.rows] == ["a", "b", "d", "e"]


def test_deduplicate_composite_text_columns() -> None:
    t = Table(["Ad", "İl"], [["Işık", "ANKARA"], ["ışık", "ankara"], ["Işık", "İzmir"]])
    assert deduplicate(t, ["Ad", "İl"]) == [1]
    with pytest.raises(UserError):
        deduplicate(t, ["Yok"])


def test_filter_rules() -> None:
    rows: list[list[Any]] = [["ended", "Ankara", 45], ["missed", "İzmir", None], ["", "ankara", 30]]

    def run(*rules: dict[str, Any], match: str = "all") -> list[list[Any]]:
        t = Table(["Durum", "İl", "Yaş"], [list(r) for r in rows])
        apply_rules(t, [FilterRule.model_validate(r) for r in rules], match=match)
        return t.rows

    assert len(run({"column": "İl", "op": "equals", "value": "ANKARA"})) == 2
    assert len(run({"column": "Durum", "op": "in", "value": ["ended", "(boş)"]})) == 2
    assert len(run({"column": "Durum", "op": "not_in", "value": ["ended"]})) == 2
    assert run({"column": "Yaş", "op": "empty"}) == [rows[1]]
    assert len(run({"column": "Yaş", "op": "equals", "value": 45})) == 1
    assert len(run({"column": "İl", "op": "contains", "value": "zm"})) == 1
    assert (
        len(
            run(
                {"column": "Durum", "op": "equals", "value": "missed"},
                {"column": "Yaş", "op": "equals", "value": "30"},
                match="any",
            )
        )
        == 2
    )
    with pytest.raises(ValueError):
        FilterRule.model_validate({"column": "İl", "op": "equals"})


# -- Excel safety ----------------------------------------------------------------
def test_safe_value_strips_control_characters() -> None:
    assert safe_value("a\x00b\x07c\x1f") == "abc"
    assert safe_value("satır 1\nsatır 2") == "satır 1 satır 2"
    assert safe_value("bad \ud800 surrogate") == "bad  surrogate"
    assert len(safe_value("x" * 40_000)) == 32_767
    assert safe_value(12) == 12 and safe_value(None) is None


def test_workbook_never_contains_formulas(tmp_path: Path) -> None:
    evil = ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)", "\x07=HYPERLINK(1)", "ok\x0bvalue"]
    path = tmp_path / "out.xlsx"
    write_workbook(path, Table(["=Başlık", "Değer"], [[e, i] for i, e in enumerate(evil)]), [("Özet", [], [["=x"]])])
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", cell.value
    values = [r[0].value for r in wb.worksheets[0].iter_rows(min_row=2)]
    assert values == ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)", "=HYPERLINK(1)", "okvalue"]
    # Reading it back keeps the literal text.
    assert read_xlsx(path).rows[0][0] == "=1+1"


# -- distribution math -----------------------------------------------------------
def test_sequential_and_round_robin() -> None:
    assert sequential_counts(301, 3) == [101, 100, 100]
    parts = distribute(list(range(12)), 3, "round_robin")
    assert parts == [[0, 3, 6, 9], [1, 4, 7, 10], [2, 5, 8, 11]]
    seq = distribute(list(range(10)), 3, "sequential")
    assert seq == [[0, 1, 2, 3], [4, 5, 6], [7, 8, 9]]
    assert distribute([], 2, "round_robin") == [[], []]


def test_ratio_counts_sum_to_total() -> None:
    assert ratio_counts(100, [3, 1]) == [75, 25]
    assert ratio_counts(10, [1, 1, 1]) == [4, 3, 3]
    assert ratio_counts(10, [50, 30, 20]) == [5, 3, 2]
    assert ratio_counts(2, [1, 1, 1, 1, 100]) == [0, 0, 0, 0, 2]  # v1 could go negative here
    rng = random.Random(7)
    for _ in range(500):
        total = rng.randint(0, 5000)
        weights = [rng.choice([0, rng.random() * 100]) for _ in range(rng.randint(1, 12))]
        if sum(weights) <= 0:
            weights[0] = 1
        counts = ratio_counts(total, weights)
        assert sum(counts) == total
        assert all(c >= 0 for c in counts)
        assert all(c == 0 for c, w in zip(counts, weights, strict=True) if w == 0)
        exact = [total * w / sum(weights) for w in weights]
        assert all(abs(c - x) < 1 for c, x in zip(counts, exact, strict=True))
    parts = distribute(list(range(10)), 2, "ratio", [70, 30])
    assert parts == [list(range(7)), [7, 8, 9]]


def test_check_weights() -> None:
    check_weights([50, 30, 20], 3, "percent")
    check_weights([2, 1], 2, "factor")
    for weights, n, mode in (([50, 40], 2, "percent"), ([1], 2, "factor"), ([-1, 2], 2, "factor")):
        with pytest.raises(UserError):
            check_weights(weights, n, mode)  # type: ignore[arg-type]
    with pytest.raises(UserError):
        ratio_counts(5, [0, 0])


def test_team_file_names_are_unique_and_safe() -> None:
    names = team_file_names(["a/b", "a:b", "x/b", "A_B", "CON", "..", "Ekip 1", "ekip 1"], reserved=["_ozet.xlsx"])
    assert len({n.casefold() for n in names}) == len(names)
    assert all(n.endswith(".xlsx") and "/" not in n and ":" not in n and "\\" not in n for n in names)
    assert names[0] == "a_b.xlsx" and names[1] == "a_b (2).xlsx" and names[2] == "x_b.xlsx"
    assert names[3] == "A_B (3).xlsx"
    assert names[4] != "CON.xlsx"
    assert names[7] == "ekip 1 (2).xlsx"
    assert team_file_names(["_ozet"], reserved=["_ozet.xlsx"]) == ["_ozet (2).xlsx"]


# -- operations over the API -----------------------------------------------------
def test_ops_registered_in_workflow_category(api: Api) -> None:
    catalogue = {op["name"]: op for op in api.c.get("/api/ops").json()}
    for name in ("calllog_analyze", "calllog_to_excel", "excel_filter", "excel_distribute"):
        assert catalogue[name]["category"] == "workflow"
    assert catalogue["calllog_analyze"]["returns_data"] is True
    assert catalogue["calllog_to_excel"]["max_inputs"] == 500
    assert catalogue["excel_distribute"]["accepts"] == ["xlsx"]


def test_analyze_groups_formats(api: Api) -> None:
    ids = [
        api.upload(make_call_log_pdf(RECORDS), "liste1.pdf")["id"],
        api.upload(make_table_pdf(), "tablo.pdf")["id"],
        api.upload(make_call_log_pdf(RECORDS_2), "liste2.pdf")["id"],
        api.upload(make_image_pdf(), "tarama.pdf")["id"],
        api.upload(make_pdf(1), "metin.pdf")["id"],
    ]
    job = api.run("calllog_analyze", ids)
    assert job["status"] == "done", job
    data = job["data"]
    kinds = [f["kind"] for f in data["files"]]
    assert kinds == ["call_log", "table", "call_log", "scanned", "unknown"]
    assert [f["rows"] for f in data["files"]] == [3, 2, 2, 0, 0]
    groups = {g["id"]: g for g in data["groups"]}
    assert groups["g1"]["kind"] == "call_log" and groups["g1"]["files"] == ["liste1.pdf", "liste2.pdf"]
    assert groups["g1"]["rows"] == 5 and groups["g1"]["columns"] == CALL_LOG_COLUMNS
    assert groups["g2"]["kind"] == "table" and groups["g2"]["columns"] == ["Ad", "Tutar", "Not"]
    assert not groups["g3"]["mergeable"]
    assert data["total_rows"] == 7
    assert data["files"][3]["message"]["en"].startswith("Scanned")


def test_call_logs_to_excel_with_phone_dedup(api: Api) -> None:
    a = api.upload(make_call_log_pdf(RECORDS), "liste1.pdf")
    b = api.upload(make_call_log_pdf(RECORDS_2), "liste2.pdf")
    s = api.upload(make_image_pdf(), "tarama.pdf")
    job = api.run("calllog_to_excel", [a["id"], b["id"], s["id"]], {"dedup": True})
    assert job["status"] == "done", job
    data = job["data"]
    assert data["total_rows_read"] == 5
    assert data["total_duplicates_removed"] == 1
    assert data["groups"][0]["dedup_columns"] == ["Telefon"]
    assert data["skipped"][0]["name"] == "tarama.pdf"
    xlsx = api.download(job["result"]["id"])
    rows = sheet_values(xlsx, "Birleşik Kayıtlar")
    assert rows[0] == ["Sıra", "Kaynak PDF", *CALL_LOG_COLUMNS]
    assert [r[0] for r in rows[1:]] == [1, 2, 3, 4]
    assert [r[1] for r in rows[1:]] == ["liste1.pdf", "liste1.pdf", "liste1.pdf", "liste2.pdf"]
    assert rows[4][3] == "Şükrü Öz" and rows[4][rows[0].index("İkamet ili")] == "İzmir"
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Birleşik Kayıtlar", "Özet"]
    injected = wb["Birleşik Kayıtlar"].cell(row=4, column=4)
    assert injected.value.startswith("=HYPERLINK") and injected.data_type == "s"
    summary = sheet_values(xlsx, "Özet")
    flat = [v for r in summary for v in r if v is not None]
    assert "Silinen mükerrer" in flat and "tarama.pdf" in flat


def test_mixed_groups_are_zipped_and_table_dedup(api: Api) -> None:
    grid = [["Ad", "Telefon", "Sehir"], ["Ali", "0532 000 00 01", "Ankara"], ["Veli", "0532 000 00 02", "Bursa"]]
    grid2 = [grid[0], ["Ali B.", "+90 532 000 00 01", "Ankara"], ["Can", "0532 000 00 03", "Izmir"]]
    ids = [
        api.upload(make_call_log_pdf(RECORDS), "cagri.pdf")["id"],
        api.upload(make_grid_pdf(grid), "t1.pdf")["id"],
        api.upload(make_grid_pdf(grid2, pages=2), "t2.pdf")["id"],
    ]
    job = api.run("calllog_to_excel", ids, {"dedup": True, "dedup_columns": "Telefon"})
    assert job["status"] == "done", job
    groups = job["data"]["groups"]
    assert [g["kind"] for g in groups] == ["call_log", "table"]
    assert groups[1]["rows_read"] == 4  # repeated header on page 2 is not a row
    assert groups[1]["duplicates_removed"] == 1
    with zipfile.ZipFile(io.BytesIO(api.download(job["result"]["id"]))) as zf:
        names = zf.namelist()
        assert names == ["cagri_kayitlari_3_kayit.xlsx", "tablo_g2_3_kayit.xlsx"]
        rows = sheet_values(zf.read(names[1]))
    assert rows[0] == ["Sıra", "Kaynak PDF", "Ad", "Telefon", "Sehir"]
    assert [r[2] for r in rows[1:]] == ["Ali", "Veli", "Can"]


def test_table_dedup_without_known_columns_fails(api: Api) -> None:
    f = api.upload(make_table_pdf(), "tablo.pdf")
    job = api.run("calllog_to_excel", [f["id"]], {"dedup": True, "dedup_columns": ["Yok"]})
    assert job["status"] == "failed"
    assert "column" in job["error"]["en"]


def test_nothing_to_merge_is_user_error(api: Api) -> None:
    f = api.upload(make_image_pdf(), "tarama.pdf")
    job = api.run("calllog_to_excel", [f["id"]])
    assert job["status"] == "failed"
    assert "Nothing to merge" in job["error"]["en"]


def test_excel_filter(api: Api) -> None:
    rows: list[list[Any]] = [
        ["Sıra", "Durum", "Telefon", "İl"],
        [1, "ended", "0532 111 22 33", "Ankara"],
        [2, "missed", "0533 111 22 33", "İzmir"],
        [3, "ended", "+90 532 111 22 33", "ankara"],
        [4, "ended", "0535 111 22 33", None],
        [5, "completed", "0536 111 22 33", "Bursa"],
    ]
    f = api.upload(make_xlsx_rows(rows), "liste.xlsx")
    rules = '[{"column": "Durum", "op": "in", "value": ["ended", "completed"]}, {"column": "İl", "op": "not_empty"}]'
    job = api.run("excel_filter", [f["id"]], {"rules": rules, "dedup_columns": "Telefon"})
    assert job["status"] == "done", job
    assert job["data"] == {
        "rows_before": 5,
        "removed_by_rules": 2,
        "duplicates_removed": 1,
        "rows_after": 2,
        "columns": ["Sıra", "Durum", "Telefon", "İl"],
    }
    out = sheet_values(api.download(job["result"]["id"]))
    assert out[1:] == [[1, "ended", "0532 111 22 33", "Ankara"], [2, "completed", "0536 111 22 33", "Bursa"]]

    bad = api.run("excel_filter", [f["id"]], {"rules": [{"column": "Yok", "op": "empty"}]})
    assert bad["status"] == "failed" and "Column not found" in bad["error"]["en"]
    assert api.start("excel_filter", [f["id"]], {"rules": "not json"}).status_code == 422
    nothing = api.run("excel_filter", [f["id"]], {})
    assert nothing["status"] == "failed"


def test_excel_distribute_ratio_and_unique_names(api: Api) -> None:
    rows: list[list[Any]] = [["Sıra", "Ad"], *[[i, f"kişi {i}"] for i in range(1, 11)]]
    f = api.upload(make_xlsx_rows(rows), "liste.xlsx")
    job = api.run(
        "excel_distribute",
        [f["id"]],
        {"teams": "a/b, a:b, Ekip 3", "mode": "ratio", "weight_mode": "percent", "weights": "50, 30, 20"},
    )
    assert job["status"] == "done", job
    assert job["data"]["counts"] == {"a/b": 5, "a:b": 3, "Ekip 3": 2}
    with zipfile.ZipFile(io.BytesIO(api.download(job["result"]["id"]))) as zf:
        names = zf.namelist()
        assert names == ["a_b.xlsx", "a_b (2).xlsx", "Ekip 3.xlsx", "_dagitim_ozeti.xlsx"]
        assert [r[0] for r in sheet_values(zf.read("a_b (2).xlsx"))[1:]] == [6, 7, 8]
        summary = sheet_values(zf.read("_dagitim_ozeti.xlsx"))
    assert summary[0] == ["Ekip", "Dosya", "Satır", "Pay (%)", "Oran"]
    assert summary[1][:3] == ["a/b", "a_b.xlsx", 5]


def test_excel_distribute_round_robin_and_validation(api: Api) -> None:
    rows: list[list[Any]] = [["Sıra", "Ad"], *[[i, f"k{i}"] for i in range(1, 8)]]
    f = api.upload(make_xlsx_rows(rows), "liste.xlsx")
    job = api.run("excel_distribute", [f["id"]], {"teams": ["A", "B", "C"], "mode": "round_robin"})
    assert job["status"] == "done", job
    with zipfile.ZipFile(io.BytesIO(api.download(job["result"]["id"]))) as zf:
        assert [r[0] for r in sheet_values(zf.read("A.xlsx"))[1:]] == [1, 4, 7]
        assert [r[0] for r in sheet_values(zf.read("C.xlsx"))[1:]] == [3, 6]

    pct = api.run("excel_distribute", [f["id"]], {"teams": ["A", "B"], "mode": "ratio", "weights": [60, 30]})
    assert pct["status"] == "failed" and "100%" in pct["error"]["en"]
    factor = api.run(
        "excel_distribute",
        [f["id"]],
        {"teams": ["A", "B"], "mode": "ratio", "weight_mode": "factor", "weights": [3, 1]},
    )
    assert factor["status"] == "done" and factor["data"]["counts"] == {"A": 5, "B": 2}
    dup = api.run("excel_distribute", [f["id"]], {"teams": ["Ekip", "ekip"]})
    assert dup["status"] == "failed" and "unique" in dup["error"]["en"]
    assert api.start("excel_distribute", [f["id"]], {"teams": []}).status_code == 422


def test_full_workflow_chain(api: Api) -> None:
    """PDFs → merged Excel → filter → distribute, each a separate job."""
    a = api.upload(make_call_log_pdf([*RECORDS, record(4, "Deniz", "0566 000 11 22", status="missed")]), "a.pdf")
    b = api.upload(make_call_log_pdf(RECORDS_2), "b.pdf")
    merged = api.run("calllog_to_excel", [a["id"], b["id"]], {"dedup": True})
    assert merged["status"] == "done", merged
    assert merged["data"]["total_rows_written"] == 5
    x = api.upload(api.download(merged["result"]["id"]), merged["result"]["name"])
    assert x["kind"] == "xlsx"
    filtered = api.run("excel_filter", [x["id"]], {"rules": [{"column": "Durum", "op": "equals", "value": "ended"}]})
    assert filtered["status"] == "done", filtered
    assert filtered["data"]["rows_after"] == 4
    y = api.upload(api.download(filtered["result"]["id"]), filtered["result"]["name"])
    dist = api.run("excel_distribute", [y["id"]], {"teams": "Kuzey\nGüney", "mode": "sequential"})
    assert dist["status"] == "done", dist
    assert dist["data"]["counts"] == {"Kuzey": 2, "Güney": 2}
    with zipfile.ZipFile(io.BytesIO(api.download(dist["result"]["id"]))) as zf:
        rows = sheet_values(zf.read("Güney.xlsx"))
    assert rows[0][:3] == ["Sıra", "Kaynak PDF", "Kayıt No"]
    assert [r[1] for r in rows[1:]] == ["a.pdf", "b.pdf"]
