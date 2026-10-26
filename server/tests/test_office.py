from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pymupdf
import pytest

from adminpdf.ops.office import find_soffice, to_cell

from .conftest import Api, make_docx, make_pdf, make_table_pdf, make_xlsx

needs_lo = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")


def text_of(data: bytes) -> str:
    with pymupdf.open("pdf", data) as doc:
        return "\n".join(p.get_text() for p in doc)


@pytest.mark.parametrize(
    ("raw", "mode", "expected"),
    [
        ("1.250,50", "tr", 1250.5),
        ("1,250.50", "en", 1250.5),
        ("300", "tr", 300),
        ("12.5", "tr", "12.5"),  # ambiguous in TR → kept as text
        ("05321234567", "tr", 5321234567),
        ("abc\x01", "tr", "abc"),
        ("1.250,50", "text", "1.250,50"),
    ],
)
def test_to_cell(raw: str, mode: str, expected: object) -> None:
    assert to_cell(raw, mode) == expected


def test_pdf_to_excel_tables_and_no_formula_injection(api: Api) -> None:
    from openpyxl import load_workbook

    f = api.upload(make_table_pdf())
    job = api.run("pdf_to_excel", [f["id"]])
    assert job["status"] == "done", job
    wb = load_workbook(io.BytesIO(api.download(job["result"]["id"])))
    ws = wb.worksheets[0]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0] == ["Ad", "Tutar", "Not"]
    assert values[1][1] == 1250.5
    assert ws.cell(row=2, column=3).data_type == "s"  # literal text, not a formula


def test_pdf_to_excel_without_tables_is_user_error(api: Api) -> None:
    f = api.upload(make_pdf(1))
    job = api.run("pdf_to_excel", [f["id"]])
    assert job["status"] == "failed"
    assert "No tables" in job["error"]["en"]


def test_pdf_to_word(api: Api) -> None:
    import docx

    f = api.upload(make_pdf(2, text="Paragraf"))
    job = api.run("pdf_to_word", [f["id"]])
    assert job["status"] == "done", job
    d = docx.Document(io.BytesIO(api.download(job["result"]["id"])))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Paragraf 1" in text and "Paragraf 2" in text


def test_office_to_pdf_builtin_engine(api: Api) -> None:
    d = api.upload(make_docx(), "rapor.docx")
    x = api.upload(make_xlsx(), "tablo.xlsx")
    assert d["kind"] == "docx" and x["kind"] == "xlsx"
    job = api.run("office_to_pdf", [d["id"], x["id"]], {"engine": "builtin"})
    assert job["status"] == "done", job
    with zipfile.ZipFile(io.BytesIO(api.download(job["result"]["id"]))) as zf:
        word = text_of(zf.read("rapor.pdf"))
        excel = text_of(zf.read("tablo.pdf"))
    assert "İĞŞÇÖÜ" in word and "Şükrü" in word
    assert "Satış" in excel and "Çay" in excel


@needs_lo
def test_office_to_pdf_libreoffice(api: Api) -> None:
    d = api.upload(make_docx(), "rapor.docx")
    job = api.run("office_to_pdf", [d["id"]], {"engine": "libreoffice"})
    assert job["status"] == "done", job
    assert "Şükrü" in text_of(api.download(job["result"]["id"]))


def test_zip_bomb_rejected(api: Api) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<x/>")
        zf.writestr("word/document.xml", b"0" * (50 * 1024 * 1024))
    f = api.upload(buf.getvalue(), "bomb.docx")
    job = api.run("office_to_pdf", [f["id"]], {"engine": "builtin"})
    assert job["status"] == "failed"
    assert "zip bomb" in job["error"]["en"]


@needs_lo
def test_libreoffice_cannot_read_local_files_or_the_network(tmp_path: Path) -> None:
    """A crafted workbook asks LibreOffice to fetch a URL (SSRF) and a Word
    field to include a local file; neither may reach the output PDF."""
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    from openpyxl import Workbook

    from adminpdf.ops.base import InputFile
    from adminpdf.ops.office import libreoffice_to_pdf

    hits: list[str] = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            hits.append(self.path)
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"SIZINTI-HTTP")

        def log_message(self, *_args: object) -> None:
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        secret = tmp_path / "secret.txt"
        secret.write_text("SIZINTI-DOSYA")
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = f'=_xlfn.WEBSERVICE("http://127.0.0.1:{server.server_port}/x")'
        ws["A2"] = f'=_xlfn.WEBSERVICE("{secret.as_uri()}")'
        wb.save(tmp_path / "evil.xlsx")

        base = io.BytesIO(make_docx("PLACEHOLDER"))
        evil_docx = tmp_path / "evil.docx"
        with zipfile.ZipFile(base) as src, zipfile.ZipFile(evil_docx, "w") as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "word/document.xml":
                    field = (
                        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve">'
                        f' INCLUDETEXT "{secret}" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/>'
                        '</w:r><w:r><w:t>cached</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>'
                    )
                    run = "<w:r><w:rPr><w:b/></w:rPr><w:t>PLACEHOLDER</w:t></w:r>"
                    data = data.decode().replace(run, field).encode()
                    assert b"INCLUDETEXT" in data
                dst.writestr(info, data)

        for name in ("evil.xlsx", "evil.docx"):
            out = tmp_path / f"{name}.pdf"
            libreoffice_to_pdf(find_soffice() or "", InputFile(tmp_path / name, name), out)
            with pymupdf.open(out) as doc:
                text = "".join(p.get_text() for p in doc)
            assert "SIZINTI" not in text, name
        assert hits == []
    finally:
        server.shutdown()


def test_libreoffice_profile_is_hardened(tmp_path: Path) -> None:
    import xml.etree.ElementTree as ET

    from adminpdf.ops.office import write_hardened_profile

    write_hardened_profile(tmp_path / "p")
    root = ET.parse(tmp_path / "p" / "user" / "registrymodifications.xcu").getroot()
    oor = "{http://openoffice.org/2001/registry}"
    values = {
        (item.get(f"{oor}path"), prop.get(f"{oor}name")): prop.findtext("value") for item in root for prop in item
    }
    assert values[("/org.openoffice.Office.Common/Security/Scripting", "DisableMacrosExecution")] == "true"
    assert values[("/org.openoffice.Office.Common/Security/Scripting", "DisableActiveContent")] == "true"
    assert values[("/org.openoffice.Office.Calc/Content/Update", "Link")] == "1"
    assert values[("/org.openoffice.Inet/Settings", "ooInetHTTPProxyPort")] == "9"
